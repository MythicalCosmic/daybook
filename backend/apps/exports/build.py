"""Building export files (xlsx with openpyxl, csv with a BOM so Excel reads Cyrillic) and
sending them to Telegram when the requester opted in.

A plain module: `exports.tasks` and `exports.services` both use it, and it imports neither.
The requester's permission is checked again here, since time passes between asking and building.
"""

import csv
import io
import logging
import zoneinfo
from datetime import date, datetime
from decimal import Decimal

from django.core.files.base import ContentFile
from django.db import transaction
from django.utils import timezone, translation
from django.utils.text import slugify
from django.utils.translation import gettext as _
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from apps.console import pulse
from apps.exports.models import Export
from apps.exports.registry import Column, Exporter, get_exporter, may_export
from apps.inbox.delivery import bot_for
from apps.members.models import Membership
from apps.notifications.tasks import send_telegram_document

logger = logging.getLogger(__name__)

# Beyond this an export is a database dump, not a spreadsheet; narrow the filters.
MAX_ROWS = 100_000
# Text that Excel or LibreOffice would run as a formula when the file is opened.
FORMULA_START = ("=", "+", "-", "@", "\t", "\r")
SHEET_TITLE_FORBIDDEN = str.maketrans({c: " " for c in "[]:*?/\\"})
EXPORTS_PULSE = "exports.exports"


class ExportFailed(Exception):
    """Building stopped; the message is safe to show to the requester."""


def validated_params(exporter: Exporter, params: dict) -> dict:
    """The exporter's params serializer applied to `params` (raises DRF's ValidationError)."""
    if exporter.params is None:
        return dict(params)
    serializer = exporter.params(data=params)
    serializer.is_valid(raise_exception=True)
    return dict(serializer.validated_data)


def build(export_id: int) -> None:
    """Build one pending export. Safe to run twice: only a pending export is built."""
    with transaction.atomic():
        export = (
            Export.objects.select_for_update(of=("self",))
            .select_related("requested_by__user", "requested_by__role", "company")
            .filter(pk=export_id)
            .first()
        )
        if export is None or export.status != Export.Status.PENDING:
            return
        export.status = Export.Status.RUNNING
        export.save(update_fields=["status", "updated_at"])

    member = export.requested_by
    with translation.override(member.user.language):
        try:
            exporter = _check(export, member)
            params = validated_params(exporter, export.params)
            content, count = _write(export, exporter, exporter.rows(member, params))
        except ExportFailed as exc:
            _finish(export, error=str(exc))
            return
        except Exception:
            logger.exception("Export %s (%s) failed", export.pk, export.kind)
            _finish(export, error=_("Something went wrong while building this export."))
            return
        export.file.save(f"{export.kind}.{export.format}", ContentFile(content), save=False)
        _finish(export, rows=count)
        if export.deliver_telegram:
            _send_to_telegram(export, exporter, member)


def download_name(export: Export, exporter: Exporter | None = None) -> str:
    """What the file is called for the person: `tasks-2026-10-20.xlsx`."""
    exporter = exporter or get_exporter(export.kind)
    label = slugify(str(exporter.label)) if exporter else ""
    day = timezone.localtime(export.created_at, _tz(export)).date().isoformat()
    return f"{label or slugify(export.kind) or 'export'}-{day}.{export.format}"


def _check(export: Export, member: Membership) -> Exporter:
    exporter = get_exporter(export.kind)
    if exporter is None:
        raise ExportFailed(_("This export is no longer available."))
    member.refresh_from_db(fields=["status", "role"])
    if (
        member.status != Membership.Status.APPROVED
        or not export.company.is_active
        or not may_export(member, exporter)
    ):
        raise ExportFailed(_("You no longer have permission to export this."))
    if export.format not in exporter.formats:
        raise ExportFailed(_("This export can't be made in that format."))
    return exporter


def _write(export: Export, exporter: Exporter, rows) -> tuple[bytes, int]:
    tz = _tz(export)
    if export.format == Export.Format.XLSX:
        return _xlsx(exporter, rows, tz)
    if export.format == Export.Format.CSV:
        return _csv(exporter, rows, tz)
    raise ExportFailed(_("This export can't be made in that format."))


def _xlsx(exporter: Exporter, rows, tz) -> tuple[bytes, int]:
    book = Workbook(write_only=True)
    sheet = book.create_sheet(title=_sheet_title(str(exporter.label)))
    sheet.freeze_panes = "A2"
    for index, column in enumerate(exporter.columns, start=1):
        width = min(max(len(str(column.label)) + 2, 12), 50)
        sheet.column_dimensions[get_column_letter(index)].width = width
    bold = Font(bold=True)
    header = []
    for column in exporter.columns:
        cell = WriteOnlyCell(sheet, value=str(column.label))
        cell.font = bold
        header.append(cell)
    sheet.append(header)
    count = 0
    for row in rows:
        count = _count(count)
        sheet.append([_xlsx_cell(sheet, col, row.get(col.key), tz) for col in exporter.columns])
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue(), count


def _xlsx_cell(sheet, column: Column, value, tz) -> WriteOnlyCell:
    if value is None:
        return WriteOnlyCell(sheet, value=None)
    match column.kind:
        case "money":
            cell = WriteOnlyCell(sheet, value=Decimal(value))
            cell.number_format = "#,##0.00"
        case "number":
            number = value if isinstance(value, int | float) else Decimal(value)
            cell = WriteOnlyCell(sheet, value=number)
        case "date":
            cell = WriteOnlyCell(sheet, value=_as_date(value))
            cell.number_format = "yyyy-mm-dd"
        case "datetime":
            cell = WriteOnlyCell(sheet, value=_local(value, tz))
            cell.number_format = "yyyy-mm-dd hh:mm"
        case "bool":
            cell = WriteOnlyCell(sheet, value=bool(value))
        case _:
            cell = WriteOnlyCell(sheet, value=str(value))
            cell.data_type = "s"  # never a formula, whatever the text starts with
    return cell


def _csv(exporter: Exporter, rows, tz) -> tuple[bytes, int]:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow([str(column.label) for column in exporter.columns])
    count = 0
    for row in rows:
        count = _count(count)
        writer.writerow([_csv_value(col, row.get(col.key), tz) for col in exporter.columns])
    # utf-8-sig: the BOM makes Excel read o'zbek and Cyrillic text correctly.
    return buffer.getvalue().encode("utf-8-sig"), count


def _csv_value(column: Column, value, tz) -> str:
    if value is None:
        return ""
    match column.kind:
        case "money":
            return f"{Decimal(value):.2f}"
        case "number":
            return str(value)
        case "date":
            return _as_date(value).isoformat()
        case "datetime":
            return _local(value, tz).strftime("%Y-%m-%d %H:%M")
        case "bool":
            return _("Yes") if value else _("No")
    text = str(value)
    return f"'{text}" if text.startswith(FORMULA_START) else text


def _count(count: int) -> int:
    count += 1
    if count > MAX_ROWS:
        raise ExportFailed(
            _("More than %(max)s rows; narrow the filters and try again.") % {"max": MAX_ROWS}
        )
    return count


def _as_date(value) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return date.fromisoformat(str(value))


def _local(value, tz) -> datetime:
    """A moment as the company's wall-clock time (spreadsheets have no time zones)."""
    if not isinstance(value, datetime):
        value = datetime.fromisoformat(str(value))
    if timezone.is_aware(value):
        value = value.astimezone(tz)
    return value.replace(tzinfo=None)


def _tz(export: Export):
    return zoneinfo.ZoneInfo(export.company.timezone)


def _sheet_title(label: str) -> str:
    return label.translate(SHEET_TITLE_FORBIDDEN).strip()[:31] or "Export"


def _finish(export: Export, *, rows: int | None = None, error: str = "") -> None:
    export.status = Export.Status.FAILED if error else Export.Status.READY
    export.rows = rows
    export.error = error[:300]
    export.finished_at = timezone.now()
    export.save(update_fields=["status", "rows", "error", "finished_at", "file", "updated_at"])
    pulse.bump(export.company_id, EXPORTS_PULSE)


def _send_to_telegram(export: Export, exporter: Exporter, member: Membership) -> None:
    user = member.user
    if user.telegram_id is None:
        return
    caption = _("%(label)s · %(count)s rows") % {"label": exporter.label, "count": export.rows}
    send_telegram_document.delay(
        bot_for(user, "owner"),
        user.telegram_id,
        export.file.name,
        storage="private",
        caption=caption,
        filename=download_name(export, exporter),
    )
