"""Finance exports: ledger entries, P&L, cash flow and debts, as Excel opens them (CSV with a
UTF-8 BOM), with the screens' filters, branch scope and hidden salary detail."""

import csv
import io
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from apps.core.storage import private_storage
from apps.exports import build
from apps.exports import services as export_services
from apps.finance import queries, services
from apps.finance.tests.factories import CounterpartyFactory
from apps.members.permissions import Perm

pytestmark = pytest.mark.django_db
D = Decimal


def built(member, kind, format="csv", params=None):
    export = export_services.request_export(member, kind, format=format, params=params or {})
    build.build(export.pk)
    export.refresh_from_db()
    assert export.status == "ready", export.error
    return export


def csv_rows(export) -> list[list[str]]:
    with private_storage.open(export.file.name, "rb") as file:
        raw = file.read()
    assert raw.startswith(b"\xef\xbb\xbf")  # Excel reads o‘zbek and Cyrillic text right
    return list(csv.reader(io.StringIO(raw.decode("utf-8-sig"))))


@pytest.fixture
def today(fin):
    return queries.local_today(fin)


@pytest.fixture
def boss(fin, member):
    return member(fin, "Owner")


def test_ledger_entries_follow_the_filters_and_hide_salaries(
    fin, boss, cash, today, member, role, person
):
    worker = member(fin, "Employee", first_name="Madina")
    services.record_transaction(
        fin, by=boss, kind="opening", amount=D("9000000"), occurred_on=today, account=cash
    )
    services.record_transaction(
        fin,
        by=boss,
        kind="expense",
        amount=D("4600000"),
        occurred_on=today,
        account=cash,
        category="salaries",
        membership=worker,
        description="Madina, oktabr",
    )
    services.record_transaction(
        fin,
        by=boss,
        kind="expense",
        amount=D("300000"),
        occurred_on=today,
        account=cash,
        category="rent",
        description="Do‘kon ijarasi",
    )
    rows = csv_rows(built(boss, "finance.transactions", params={"kind": ["expense"]}))
    assert rows[0][:3] == ["Date", "Kind", "Status"]
    assert {row[8] for row in rows[1:]} == {"Madina, oktabr", "Do‘kon ijarasi"}
    viewer = person(
        fin, role(fin, "Books", [Perm.FINANCE_VIEW, Perm.BRANCHES_ALL, Perm.DATA_EXPORT])
    )
    rows = csv_rows(built(viewer, "finance.transactions", params={"kind": "expense"}))
    salary = next(row for row in rows[1:] if row[9] == "-4600000.00")
    assert (salary[6], salary[8]) == ("", "")


def test_the_pnl_export_has_a_row_per_period_and_category(fin, boss, cash, today):
    services.record_transaction(
        fin,
        by=boss,
        kind="income",
        amount=D("700"),
        occurred_on=today,
        account=cash,
        category="sales",
    )
    services.record_transaction(
        fin,
        by=boss,
        kind="expense",
        amount=D("200"),
        occurred_on=today,
        account=cash,
        category="rent",
    )
    export = built(boss, "finance.pnl", format="xlsx", params={"preset": "month"})
    with private_storage.open(export.file.name, "rb") as file:
        sheet = load_workbook(file).active
    values = [[cell.value for cell in row] for row in sheet.iter_rows(min_row=2)]
    assert [row[1] for row in values] == ["Income", "Expenses", "Profit"]
    assert values[-1][4] == D("500")


def test_cash_flow_and_debts_exports(fin, boss, cash, today):
    services.record_transaction(
        fin,
        by=boss,
        kind="income",
        amount=D("700"),
        occurred_on=today,
        account=cash,
        category="sales",
    )
    rows = csv_rows(built(boss, "finance.cash_flow", params={"preset": "today"}))
    assert rows[1][0] == "Opening cash"
    assert rows[-1] == ["Closing cash", "", "", "700.00"]
    services.open_debt(
        fin,
        by=boss,
        direction="receivable",
        kind="trade",
        amount=D("150"),
        currency="UZS",
        issued_on=today,
        counterparty=CounterpartyFactory(company=fin, name="Sardor"),
        is_opening=True,
    )
    rows = csv_rows(built(boss, "finance.debts", params={"direction": "receivable"}))
    assert rows[1][2] == "Sardor"
    assert rows[1][8] == "150.00"


def test_ledger_and_debt_exports_are_marked_sensitive(fin, boss):
    assert export_services.request_export(boss, "finance.transactions", format="csv").sensitive
    assert export_services.request_export(boss, "finance.debts", format="csv").sensitive
    assert not export_services.request_export(boss, "finance.pnl", format="csv").sensitive
